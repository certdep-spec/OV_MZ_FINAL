"""
CertifyPro - Точка входу з екраном вибору клієнта
v3.0: Мультиклієнтська архітектура
"""

import logging
import os
import sys
from pathlib import Path

logger = logging.getLogger(__name__)

# Явні іморти для PyInstaller
import yaml  # noqa: F401

try:
    import openpyxl  # noqa: F401
    import pandas  # noqa: F401
except ImportError:
    pass

# ===== Додаємо core/ в PYTHONPATH для коректних імпортів =====
CORE_DIR = Path(__file__).parent
if str(CORE_DIR) not in sys.path:
    sys.path.insert(0, str(CORE_DIR))

# Примусово UTF-8 для консолі Windows
if sys.platform == "win32":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")
    os.environ["PYTHONIOENCODING"] = "utf-8"


def select_client_gui():
    """Показати GUI діалог вибору клієнта."""
    import customtkinter as ctk

    ctk.set_appearance_mode("dark")
    ctk.set_default_color_theme("blue")

    root = ctk.CTk()
    root.title("CertifyPro - Вибір клієнта")
    root.geometry("700x480")
    root.resizable(False, False)

    # Центрування вікна
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    x = (screen_width - 700) // 2
    y = (screen_height - 480) // 2
    root.geometry(f"700x480+{x}+{y}")

    selected_client = None

    def on_select(client_id):
        nonlocal selected_client
        selected_client = client_id
        root.destroy()

    # Заголовок
    title = ctk.CTkLabel(
        root,
        text="🚀 CertifyPro v3.0",
        font=ctk.CTkFont(size=24, weight="bold"),
    )
    title.pack(pady=(30, 5))

    subtitle = ctk.CTkLabel(
        root,
        text="Оберіть клієнта для продовження:",
        font=ctk.CTkFont(size=14),
        text_color="gray",
    )
    subtitle.pack(pady=(0, 20))

    # Кнопки клієнтів
    from client_manager import ClientManager
    from config import PROFILES_PATH

    manager = ClientManager()
    manager.load_profiles(PROFILES_PATH)

    for client_id, profile in manager.get_all_profiles().items():
        # Frame для кожного клієнта
        frame = ctk.CTkFrame(root)
        frame.pack(pady=8, padx=50, fill="x")

        # Ліва частина - текст (займає все доступне місце мінус кнопка)
        text_frame = ctk.CTkFrame(frame, fg_color="transparent")
        text_frame.pack(side="left", fill="x", expand=True, padx=15, pady=12)

        # Назва компанії (жирний текст, один рядок)
        company_label = ctk.CTkLabel(
            text_frame,
            text=profile.display_name,
            font=ctk.CTkFont(size=13, weight="bold"),
            anchor="w",
            justify="left",
        )
        company_label.pack(anchor="w")

        # Опис (звичайний текст, менший розмір)
        desc_label = ctk.CTkLabel(
            text_frame,
            text=profile.description,
            font=ctk.CTkFont(size=11),
            text_color="gray60",
            anchor="w",
            justify="left",
            wraplength=400,
        )
        desc_label.pack(anchor="w")

        # Кнопка вибору (фіксована ширина, праворуч, не стискається)
        btn = ctk.CTkButton(
            frame,
            text="Обрати →",
            command=lambda cid=client_id: on_select(cid),
            width=100,
            height=35,
            fg_color="#2B7CB9",
            hover_color="#1A5F8A",
        )
        btn.pack(side="right", padx=15, pady=12)

    # Кнопка виходу
    exit_btn = ctk.CTkButton(
        root,
        text="Вихід",
        command=root.destroy,
        fg_color="gray",
        hover_color="#555555",
        width=150,
    )
    exit_btn.pack(pady=20)

    root.mainloop()
    return selected_client


def main():
    """Головна функція запуску з обробкою помилок."""
    try:
        if len(sys.argv) > 1:
            client_id = sys.argv[1]
            logger.info(f"🚀 Запуск CertifyPro для клієнта: {client_id} (з аргументу)")
            logger.debug(f"📋 sys.argv = {sys.argv}")
        else:
            # Показуємо екран вибору клієнта
            client_id = select_client_gui()
            if client_id is None:
                logger.warning("❌ Клієнта не обрано. Вихід.")
                sys.exit(0)
            logger.info(f"🚀 Запуск CertifyPro для клієнта: {client_id}")

        logger.debug("🔍 Перед ініціалізацією config...")

        # Ініціалізуємо конфігурацію клієнта
        import config

        config.initialize_app()
        config.init_client_config(client_id)

        # Імпортуємо GUI після ініціалізації конфігурації
        from gui.main_window import MainWindow

        app = MainWindow()
        app.mainloop()

    except Exception as e:
        import traceback

        error_msg = f"Критична помилка при запуску:\n\n{e}\n\n{traceback.format_exc()}"
        logger.critical(error_msg)

        # Показуємо GUI помилку, бо консоль може закритися
        try:
            import tkinter as tk
            from tkinter import messagebox

            root = tk.Tk()
            root.withdraw()
            messagebox.showerror("CertifyPro: Критична помилка", error_msg)
            root.destroy()
        except Exception:
            logger.error(error_msg)
        sys.exit(1)


if __name__ == "__main__":
    main()
