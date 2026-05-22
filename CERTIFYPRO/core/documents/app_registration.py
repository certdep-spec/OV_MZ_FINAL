"""
Реєстрація заявки — заміна placeholder'ів у шаблоні заявки.

Виділено з ZayavkaGenerator для дотримання SRP.
"""

import logging
from pathlib import Path
from typing import Dict

from docx import Document

from documents.document_base import DocumentBase

logger = logging.getLogger(__name__)


class ApplicationRegistrar:
    """
    Автоматична реєстрація заявки.

    Замінює placeholder'и реєстраційним номером та датою
    у вже згенерованому документі заявки.
    """

    @staticmethod
    def register(doc_path: Path, app_data: Dict) -> bool:
        """
        Замінити placeholder'и реєстраційними даними.

        Args:
            doc_path: Шлях до документу заявки
            app_data: Дані заявки (app_number, app_date)

        Returns:
            True якщо успішно
        """
        try:
            reg_data = DocumentBase.format_registration_data(
                app_data.get("app_number", ""), app_data.get("app_date", "")
            )
            reg_number = reg_data["reg_number"]
            reg_date = reg_data["reg_date"]

            doc = Document(doc_path)

            def process_paragraphs(paragraphs):
                for paragraph in paragraphs:
                    if _needs_registration_replacement(paragraph.text):
                        for run in paragraph.runs:
                            if _needs_registration_replacement(run.text):
                                run.text = (
                                    run.text.replace("114/____ТРП", reg_number)
                                    .replace("114/ ТРП", reg_number)
                                    .replace("114/ТРП", reg_number)
                                )
                    if _needs_date_replacement(paragraph.text):
                        for run in paragraph.runs:
                            if _needs_date_replacement(run.text):
                                run.text = run.text.replace(
                                    "«____» ______________202_ р.", reg_date
                                ).replace("«____» __________ 202_ р.", reg_date)

            # Обробка основного тексту
            process_paragraphs(doc.paragraphs)

            # Обробка колонтитулів
            for section in doc.sections:
                process_paragraphs(section.footer.paragraphs)
                process_paragraphs(section.header.paragraphs)
                for table in section.footer.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            process_paragraphs(cell.paragraphs)

            # Обробка таблиць
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        process_paragraphs(cell.paragraphs)

            doc.save(doc_path)
            logger.info(f"✅ Заявку зареєстровано: {reg_number}, {reg_date}")
            return True

        except Exception as e:
            logger.error(f"Помилка реєстрації заявки: {e}", exc_info=True)
            return False


def _needs_registration_replacement(text: str) -> bool:
    """Перевірити чи параграф містить placeholder реєстраційного номеру."""
    return any(p in text for p in ("114/____ТРП", "114/ ТРП", "114/ТРП"))


def _needs_date_replacement(text: str) -> bool:
    """Перевірити чи параграф містить placeholder дати."""
    return any(
        p in text for p in ("«____» ______________202_ р.", "«____» __________ 202_ р.")
    )
