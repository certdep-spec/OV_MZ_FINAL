"""
Побудова додатків до заявки — групування партій за продуктами.

Виділено з ZayavkaGenerator для дотримання SRP.
"""

import logging
from pathlib import Path
from typing import Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

logger = logging.getLogger(__name__)

# Ширини колонок таблиці партій
_PARTY_COL_WIDTHS = [Cm(0.5), Cm(8.1), Cm(3.7), Cm(2.1), Cm(2.1)]
_INGREDIENT_COL_WIDTHS = [Cm(0.5), Cm(3.9), Cm(4.0), Cm(3.9), Cm(4.2)]


def append_appendix(
    doc_path: Path,
    app_data: Dict,
    parties: List[Dict],
    data_service,
) -> bool:
    """
    Додати додатки до кінця існуючого документу заявки.

    Один додаток на один унікальний продукт.

    Args:
        doc_path: Шлях до документу заявки
        app_data: Дані заявки (director_name)
        parties: Список партій
        data_service: DocumentDataService (для групування, версій, інгредієнтів)

    Returns:
        True якщо успішно
    """
    try:
        logger.info(f"📝 Додавання Додатків до: {doc_path}")
        doc = Document(doc_path)

        # Налаштовуємо лівий відступ до 2 см для всіх розділів документу
        for section in doc.sections:
            section.left_margin = Cm(2.0)

        products_dict = data_service.group_parties_by_product(parties)

        for product_idx, (product_name, product_parties) in enumerate(
            products_dict.items(), 1
        ):
            if product_idx > 1:
                doc.add_page_break()

            _add_product_header(doc, product_idx, product_name)
            _add_parties_table(doc, product_parties)

            selected_version = data_service.select_version_for_product(
                product_name, product_parties
            )
            ingredients = data_service.get_ingredients_with_version(
                product_name, selected_version
            )
            _add_ingredients_section(doc, ingredients)

            _add_director_signature(doc, app_data)

        doc.save(doc_path)
        logger.info(f"✅ Додатки додано до Заявки: {doc_path}")
        return True

    except Exception as e:
        logger.error(f"Помилка додавання Додатків: {e}", exc_info=True)
        return False


def _add_product_header(doc: Document, idx: int, product_name: str) -> None:
    """Заголовок «Додаток X» + назва продукції."""
    for text, size, bold in [
        (f"Додаток {idx}", 12, True),
        ("до заявки на проведення робіт з оцінки відповідності", 12, True),
    ]:
        p = doc.add_paragraph()
        run = p.add_run(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.name = "Times New Roman"
        run.font.size = Pt(size)
        run.font.bold = bold

    p = doc.add_paragraph()
    run = p.add_run(product_name)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = False

    doc.add_paragraph()


def _add_parties_table(doc: Document, parties: List[Dict]) -> None:
    """Таблиця партій продукту."""
    p = doc.add_paragraph()
    run = p.add_run("Партії продукції:")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    doc.add_paragraph()

    tbl = doc.add_table(rows=1, cols=5)
    _set_table_borders(tbl)
    tbl.autofit = False
    _set_col_widths(tbl, _PARTY_COL_WIDTHS)

    headers = [
        "№",
        "Назва продукції",
        "Номер партії / дата",
        "Кількість",
        "Од. вим.",
    ]
    _add_table_row(tbl.rows[0], headers, bold=True, font_size=10)

    for i, party in enumerate(parties, 1):
        row = tbl.add_row()
        product_name = party.get("product_name", "")
        values = [
            str(i),
            product_name,
            f"№ {party.get('party_code', '')} від {party.get('party_date', '')}",
            str(party.get("quantity", 0)),
            party.get("unit", "шт"),
        ]
        _add_table_row(row, values, bold=False, font_size=10)
        _set_col_widths(tbl, _PARTY_COL_WIDTHS)

    doc.add_paragraph()
    doc.add_paragraph()


def _add_ingredients_section(doc: Document, ingredients: List[Dict]) -> None:
    """Секція складу продукції (ПАР)."""
    p = doc.add_paragraph()
    run = p.add_run("Склад продукції (ПАР):")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    doc.add_paragraph()

    if not ingredients:
        p = doc.add_paragraph()
        run = p.add_run("не заповнено")
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        return

    tbl = doc.add_table(rows=1, cols=5)
    _set_table_borders(tbl)
    tbl.autofit = False
    _set_col_widths(tbl, _INGREDIENT_COL_WIDTHS)

    headers = [
        "№",
        "Хімічне найменування",
        "Торговельна марка",
        "CAS №",
        "Сертифікат (№ та дата)",
    ]
    _add_table_row(tbl.rows[0], headers, bold=True, font_size=9)

    for i, ing in enumerate(ingredients, 1):
        row = tbl.add_row()
        values = [
            str(i),
            ing["chemical_name"],
            ing["trade_mark"],
            ing["cas_number"],
            ing["certificate"],
        ]
        _add_table_row(row, values, bold=False, font_size=9)
        _set_col_widths(tbl, _INGREDIENT_COL_WIDTHS)

    doc.add_paragraph()
    doc.add_paragraph()


def _add_director_signature(doc: Document, app_data: Dict) -> None:
    """Підпис директора."""
    director_name = app_data.get("director_name", "")
    sign_para = doc.add_paragraph()
    sign_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sign_para.add_run(f"Директор\t\t\t{director_name}")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def _add_table_row(row, values: List[str], bold: bool, font_size: int) -> None:
    """Заповнити рядок таблиці."""
    alignments = [
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.LEFT,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.LEFT,
    ]
    for i, (cell, value) in enumerate(zip(row.cells, values)):
        cell.text = value
        for paragraph in cell.paragraphs:
            paragraph.alignment = (
                alignments[i] if i < len(alignments) else WD_ALIGN_PARAGRAPH.LEFT
            )
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(font_size)
                run.font.bold = bold


def _set_col_widths(table, widths) -> None:
    """Встановити ширину колонок."""
    for idx, width in enumerate(widths):
        if idx < len(table.columns):
            table.columns[idx].width = width
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = width


def _set_table_borders(table) -> None:
    """Безпечно встановити межі таблиці (з обходом відсутності стилю 'Table Grid')"""
    try:
        table.style = 'Table Grid'
    except Exception:
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            
            tblPr = table._tbl.tblPr
            tblBorders = OxmlElement('w:tblBorders')
            
            for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'single')
                border.set(qn('w:sz'), '4')  # 0.5 pt
                border.set(qn('w:space'), '0')
                border.set(qn('w:color'), 'auto')
                tblBorders.append(border)
                
            tblPr.append(tblBorders)
        except Exception as e:
            logger.warning(f"Не вдалося встановити межі через XML: {e}")
