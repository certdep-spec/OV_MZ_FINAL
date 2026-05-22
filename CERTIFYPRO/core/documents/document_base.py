"""
Базовий абстрактний клас для генераторів документів
Виносить спільну логіку з vc_generator.py
"""

import calendar
import logging
import shutil
from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from docxtpl import DocxTemplate

from documents import document_utils as utils
from documents import tags as T
from exceptions import (
    DocumentGenerationError,
    TemplateNotFoundError,
    ValidationError,
)

logger = logging.getLogger(__name__)


class DocumentBase(ABC):
    """
    Абстрактний базовий клас для генераторів документів.

    Відповідає за:
    - Форматування даних (дати, номери заявок)
    - Генерацію документів через docxtpl
    - Заміну тегів у існуючих документах
    - Роботу з таблицями
    - Логування
    """

    def __init__(self, template_dir: Path, output_dir: Path):
        """
        Ініціалізація базового класу.

        Args:
            template_dir: Шлях до папки з шаблонами
            output_dir: Шлях до папки для виводу документів
        """
        self.template_dir = template_dir
        self.output_dir = output_dir

    @abstractmethod
    def get_template_name(self) -> str:
        """
        Повернути ім'я файлу шаблону.

        Returns:
            Назва файлу шаблону (наприклад, "1_Заявка_бланк_Д.docx")
        """

    @abstractmethod
    def generate(
        self,
        app_data: Dict,
        party_data: Optional[Dict] = None,
        party_id: int = 0,
        **kwargs,
    ) -> Dict[str, int]:
        """
        Генерувати документ.

        Args:
            app_data: Дані заявки (app_number, app_date, company_name, тощо)
            party_data: Дані партії (product_name, party_code, quantity, тощо)
            party_id: ID партії в БД
            **kwargs: Додаткові параметри для конкретних генераторів

        Returns:
            Словник зі статистикою (наприклад, {'zayavka': 1})
        """

    # ===== СТАТИЧНІ УТИЛІТИ =====

    @staticmethod
    def format_app_number(app_number: str) -> str:
        """Форматує номер заявки на 3 цифри з нулями попереду."""
        return utils.format_app_number(app_number)

    @staticmethod
    def create_structured_path(
        output_dir: Path, app_number: str, app_date: str
    ) -> Path:
        """
        Створення структурованої папки для документів.

        Формат: {output_dir}/{app_date}/{formatted_app_number}/

        Args:
            output_dir: Базова папка для виводу
            app_number: Номер заявки
            app_date: Дата заявки

        Returns:
            Шлях до створеної папки
        """
        formatted_num = DocumentBase.format_app_number(app_number)
        folder_path = output_dir / app_date / formatted_num
        folder_path.mkdir(parents=True, exist_ok=True)
        return folder_path

    @staticmethod
    def calculate_expiry_date(mfg_date: str, shelf_life_months: int) -> str:
        """Розрахунок терміну придатності."""
        return utils.calculate_expiry_date(mfg_date, shelf_life_months)

    @staticmethod
    def get_month_name_ua(month_num: str) -> str:
        """Отримати назву місяця українською."""
        return utils.get_month_name_ua(month_num)

    @staticmethod
    def format_registration_data(app_number: str, app_date: str) -> Dict[str, str]:
        """Формує реєстраційні дані для заявки."""
        return utils.format_registration_data(app_number, app_date)

    # ===== МЕТОДИ ГЕНЕРАЦІЇ =====

    def generate_with_docxtpl(
        self, template_name: str, context: Dict, save_path: str
    ) -> bool:
        """
        Генерація документа через docxtpl (для шаблонів з тегами {{TAG}}).

        Args:
            template_name: Назва файлу шаблону
            context: Словник з даними для заміни тегів
            save_path: Шлях для збереження результату

        Returns:
            True якщо успішно, False інакше

        Raises:
            TemplateNotFoundError: Якщо шаблон не знайдено
            DocumentGenerationError: Якщо помилка генерації
        """
        try:
            template_path = self._find_template(template_name)
            if not template_path:
                raise TemplateNotFoundError(f"Шаблон не знайдено: {template_name}")

            doc = DocxTemplate(template_path)
            doc.render(context)
            doc.save(save_path)
            logger.info(f"✅ Згенеровано (docxtpl): {save_path}")
            return True

        except TemplateNotFoundError:
            raise
        except DocumentGenerationError:
            raise
        except Exception as e:
            logger.exception(
                f"Помилка генерації документа: {template_name}",
                extra={
                    "template": (
                        str(template_path)
                        if "template_path" in locals()
                        else template_name
                    ),
                    "save_path": save_path,
                },
            )
            raise DocumentGenerationError(
                f"Помилка шаблону '{template_name}'. Перевірте правильність тегів (наприклад {{{{ТЕГ}}}} чи {{% if %}}) у файлі шаблону.\nДеталі: {e}"
            ) from e

    def generate_with_copy_and_replace(
        self, template_name: str, replacements: Dict[str, str], save_path: str
    ) -> bool:
        """
        Генерація документа через копіювання + заміну тегів (для старих шаблонів).

        Args:
            template_name: Назва файлу шаблону
            replacements: Словник {тег: значення} для заміни
            save_path: Шлях для збереження результату

        Returns:
            True якщо успішно, False інакше
        """
        try:
            template_path = self._find_template(template_name)
            if not template_path:
                raise TemplateNotFoundError(f"Шаблон не знайдено: {template_name}")

            # Копіюємо шаблон
            shutil.copy(template_path, save_path)

            # Замінюємо теги
            if self.replace_text_in_document(Path(save_path), replacements):
                logger.info(f"✅ Згенеровано (copy+replace): {save_path}")
                return True

            return False

        except (TemplateNotFoundError, DocumentGenerationError):
            raise
        except Exception as e:
            raise DocumentGenerationError(
                f"Помилка генерації {template_name}: {e}"
            ) from e

    def replace_text_in_document(
        self, doc_path: Path, replacements: Dict[str, str]
    ) -> bool:
        """
        Знайти і замінити теги [B1], {{B1}}, {B1} та інші на реальні значення.

        Працює з:
        - Основним текстом документа
        - Таблицями
        - Колонтитулами (header/footer)

        Args:
            doc_path: Шлях до документа Word
            replacements: Словник {тег: значення}

        Returns:
            True якщо заміни виконано, False інакше
        """
        try:
            doc = Document(doc_path)

            def replace_in_paragraph(paragraph, replacements):
                """Замінює теги в одному параграфі."""
                if not paragraph.runs or not paragraph.text:
                    return

                # Збираємо всі варіанти тегів ([X] → {{X}} → {X})
                all_reps = {}
                for tag, value in replacements.items():
                    if value is not None:
                        for variant in [
                            tag,
                            tag.replace("[", "{{").replace("]", "}}"),
                            tag.replace("[", "{").replace("]", "}"),
                        ]:
                            all_reps[variant] = str(value)

                full_text = paragraph.text
                if not any(t in full_text for t in all_reps):
                    return

                # Заміна кожного тегу
                for tag, value in all_reps.items():
                    idx = full_text.find(tag)
                    if idx == -1:
                        continue

                    tag_end = idx + len(tag)
                    run_ends = []
                    cumulative = 0
                    for r in paragraph.runs:
                        cumulative += len(r.text)
                        run_ends.append(cumulative)

                    first_run = None
                    last_run = None
                    for ri, end_pos in enumerate(run_ends):
                        start_pos = run_ends[ri - 1] if ri > 0 else 0
                        if start_pos < tag_end and end_pos > idx:
                            if first_run is None:
                                first_run = ri
                            last_run = ri

                    if first_run is None:
                        continue

                    combined = "".join(
                        paragraph.runs[ri].text for ri in range(first_run, last_run + 1)
                    )
                    if tag in combined:
                        combined = combined.replace(tag, value)

                    paragraph.runs[first_run].text = combined
                    for ri in range(first_run + 1, last_run + 1):
                        paragraph.runs[ri].text = ""

                    full_text = paragraph.text

            # Обробка параграфів
            for paragraph in doc.paragraphs:
                replace_in_paragraph(paragraph, replacements)

            # Обробка таблиць
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replace_in_paragraph(paragraph, replacements)

            # Обробка колонтитулів
            for section in doc.sections:
                for paragraph in section.footer.paragraphs:
                    replace_in_paragraph(paragraph, replacements)
                for paragraph in section.first_page_footer.paragraphs:
                    replace_in_paragraph(paragraph, replacements)
                for paragraph in section.header.paragraphs:
                    replace_in_paragraph(paragraph, replacements)
                for paragraph in section.first_page_header.paragraphs:
                    replace_in_paragraph(paragraph, replacements)
                try:
                    for paragraph in section.even_pages_header.paragraphs:
                        replace_in_paragraph(paragraph, replacements)
                except Exception:
                    pass
                try:
                    for paragraph in section.even_pages_footer.paragraphs:
                        replace_in_paragraph(paragraph, replacements)
                except Exception:
                    pass

            # Видалення порожніх параграфів у кінці документа (щоб уникнути зайвої сторінки)
            self._remove_trailing_empty_paragraphs(doc)

            doc.save(doc_path)
            return True

        except DocumentGenerationError:
            raise
        except Exception as e:
            raise DocumentGenerationError(f"Помилка при заміні тексту: {e}") from e

    # ===== РОБОТА З ТАБЛИЦЯМИ =====

    def set_table_borders(self, table) -> None:
        """
        Встановлює рамки для таблиці.

        Args:
            table: Об'єкт таблиці docx
        """
        tbl = table._element

        # Знаходимо або створюємо властивості таблиці (tblPr)
        tblPr = tbl.xpath("w:tblPr")
        if not tblPr:
            tblPr_elem = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr_elem)
        else:
            tblPr_elem = tblPr[0]

        # Видаляємо існуючі рамки
        old_borders = tblPr_elem.xpath("w:tblBorders")
        if old_borders:
            tblPr_elem.remove(old_borders[0])

        tblBorders = OxmlElement("w:tblBorders")

        borders = ["top", "left", "bottom", "right", "insideH", "insideV"]

        for border_name in borders:
            border = OxmlElement(f"w:{border_name}")
            border.set(qn("w:val"), "single")
            border.set(qn("w:sz"), "4")  # 4 = 1/2 pt
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "000000")
            tblBorders.append(border)

        tblPr_elem.append(tblBorders)

    def set_column_widths(self, table, col_widths: List[Cm]) -> None:
        """
        Встановлює ширину колонок таблиці.

        Args:
            table: Об'єкт таблиці docx
            col_widths: Список ширин колонок
        """
        for col_idx, width in enumerate(col_widths):
            if col_idx < len(table.columns):
                table.columns[col_idx].width = width

    # ===== ДОПОМІЖНІ МЕТОДИ =====

    def _find_template(self, template_name: str) -> Optional[Path]:
        """
        Знайти шаблон у папці templates (підтримує .docx та .doc).

        Args:
            template_name: Назва файлу шаблону

        Returns:
            Шлях до шаблону або None
        """
        template_path = self.template_dir / template_name
        if template_path.exists():
            return template_path

        # Спробувати з іншим розширенням
        if template_name.endswith(".docx"):
            alt_path = self.template_dir / template_name.replace(".docx", ".doc")
        elif template_name.endswith(".doc"):
            alt_path = self.template_dir / template_name.replace(".doc", ".docx")
        else:
            alt_path_docx = self.template_dir / f"{template_name}.docx"
            if alt_path_docx.exists():
                return alt_path_docx
            alt_path = self.template_dir / f"{template_name}.doc"

        if alt_path.exists():
            return alt_path

        return None

    def _build_base_replacements(
        self,
        app_data: Dict,
        party_data: Optional[Dict] = None,
        party_id: int = 0,
        extra: Optional[Dict] = None,
    ) -> Dict[str, str]:
        """
        Побудувати базовий словник замін для документа.

        Створює теги у форматах: [X], {{X}}, {X}

        Args:
            app_data: Дані заявки
            party_data: Дані партії (опціонально)
            party_id: ID партії
            extra: Додаткові заміни

        Returns:
            Словник {тег: значення}
        """
        formatted_app_num = self.format_app_number(app_data.get("app_number", ""))

        base = {
            T.APP_NUMBER: formatted_app_num,
            T.APP_DATE: app_data.get("app_date", ""),
            T.COMPANY_NAME: app_data.get("company_name", ""),
            T.COMPANY_CODE: app_data.get("company_code", ""),
            T.PRODUCTION_ADDRESS: app_data.get("production_address", ""),
            T.TU_CODE: app_data.get("tu_code", ""),
            T.DIRECTOR_NAME: app_data.get("director_name", ""),
        }

        if party_data:
            base[T.PRODUCT_NAME] = party_data.get("product_name", "")
            base[T.PARTY_CODE] = party_data.get("party_code", "")
            base[T.MFG_DATE] = party_data.get("mfg_date", "")

            # Розрахунок терміну придатності
            expiry_date = self.calculate_expiry_date(
                party_data.get("mfg_date", ""),
                party_data.get("shelf_life_months", 36),
            )
            base[T.EXPIRY_DATE] = expiry_date
            base[T.FIN_DATE] = expiry_date

            quantity = party_data.get("quantity", "")
            unit = party_data.get("unit", "шт")
            base[T.QUANTITY_UNIT] = f"{quantity} {unit}".strip()

            # Сертифікат та протокол
            base[T.CERT_NUMBER] = party_data.get("cert_number", "")
            base[T.CERT_DATE] = party_data.get("cert_date", "")

            protocol_number = party_data.get("protocol_number", "").strip()
            protocol_date = party_data.get("protocol_date", "").strip()
            if protocol_number and protocol_date:
                protocol_combined = f"{protocol_number} від {protocol_date} р."
            else:
                protocol_combined = protocol_number or protocol_date
            base[T.PROTOCOL] = protocol_combined

            base[T.PARTY_DATE] = party_data.get("party_date", "")
            base[T.PARTY_ID] = str(party_id)
            base[T.PARTY_ID_2] = str(party_id)
            base[T.UNIT] = party_data.get("unit", "шт")

        # Додати варіанти {{X}} та {X}
        all_replacements = {}
        for tag, value in base.items():
            if value is not None:
                all_replacements[tag] = str(value)
                # {{X}} варіант
                double_brace = tag.replace("[", "{{").replace("]", "}}")
                all_replacements[double_brace] = str(value)
                # {X} варіант
                single_brace = tag.replace("[", "{").replace("]", "}")
                all_replacements[single_brace] = str(value)

        # Додати extra заміни
        if extra:
            for tag, value in extra.items():
                if value is not None:
                    all_replacements[tag] = str(value)

        return all_replacements

    def _build_final_replacements(
        self,
        app_data: Dict,
        party_data: Dict,
        party_id: int,
        par_certificates: Optional[str] = None,
    ) -> Dict[str, str]:
        """
        Побудувати словник замін для фінальних документів
        (Сертифікат, Декларація, Угода, Рішення на видачу, Протокол аналізу, Висновок).

        Базується на _build_base_replacements() + додає [3F2] для сертифікатів ПАР.

        Args:
            app_data: Дані заявки
            party_data: Дані партії
            party_id: ID партії
            par_certificates: Сертифікати ПАР (опціонально)

        Returns:
            Словник {тег: значення} з варіантами [X], {{X}}, {X}
        """
        # Використовуємо базовий метод + додаємо специфічні теги
        replacements = self._build_base_replacements(app_data, party_data, party_id)

        # Додати сертифікати ПАР
        if par_certificates:
            replacements[T.PAR_CERTIFICATES] = str(par_certificates)
            replacements["{{3F2}}"] = str(par_certificates)
            replacements["{3F2}"] = str(par_certificates)

        return replacements

    def _remove_trailing_empty_paragraphs(self, doc) -> None:
        """
        Видаляє порожні параграфи в кінці документа, щоб уникнути зайвої порожньої сторінки.

        Залишає максимум 2 порожні параграфи для підписів.
        """
        paragraphs = doc.paragraphs
        # Знаходимо останній непорожній параграф
        last_content_idx = -1
        for i in range(len(paragraphs) - 1, -1, -1):
            p = paragraphs[i]
            text = p.text.strip() if p.text else ""
            has_image = any(
                hasattr(run, '_element') and 'drawing' in run._element.xml
                for run in p.runs
            )
            if text or has_image:
                last_content_idx = i
                break

        if last_content_idx >= 0 and last_content_idx < len(paragraphs) - 2:
            # Видаляємо порожні параграфи після останнього контентного, залишаючи 2
            max_empty = 2
            empty_count = 0
            for i in range(len(paragraphs) - 1, last_content_idx, -1):
                p = paragraphs[i]
                if not p.text.strip():
                    empty_count += 1
                    if empty_count > max_empty:
                        p_elem = p._element
                        p_elem.getparent().remove(p_elem)

    def _validate_app_data(self, app_data: Dict) -> None:
        """
        Валідація обов'язкових даних заявки.

        Args:
            app_data: Дані заявки

        Raises:
            ValidationError: Якщо обов'язкові поля відсутні
        """
        required_fields = ["app_number", "app_date", "company_name"]

        for field in required_fields:
            if field not in app_data or not app_data[field]:
                raise ValidationError(f"Обов'язкове поле '{field}' відсутнє в app_data")

    def _get_file_date(self) -> str:
        """
        Отримати поточну дату для імені файлу.

        Returns:
            Дата у форматі YYYY-MM-DD
        """
        return datetime.now().strftime("%Y-%m-%d")

    def get_output_folder(self, context) -> Path:
        """
        Отримати папку для збереження документів.

        Args:
            context: Контекст документу (DocumentContextDTO або dict з application)

        Returns:
            Path до папки для збереження
        """
        app = (
            context.application
            if hasattr(context, "application")
            else context.get("application", {})
        )
        app_number = (
            app.app_number
            if hasattr(app, "app_number")
            else app.get("app_number", "000")
        )
        app_date = (
            app.app_date
            if hasattr(app, "app_date")
            else app.get("app_date", "2026-01-01")
        )

        return self.create_structured_path(self.output_dir, app_number, app_date)
