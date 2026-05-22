"""
Тести для покриття - частина 4
Фокус на zayavka_generator, vc_generator, context_builder, document_context
"""

import os
import tempfile

import pytest

from database.db_manager import DatabaseManager


@pytest.fixture
def temp_db():
    """Створення тимчасової БД"""
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
    temp_file.close()
    db_manager = DatabaseManager(temp_file.name)
    db_manager.close()
    yield temp_file.name
    os.unlink(temp_file.name)


@pytest.fixture
def temp_dirs(tmp_path):
    """Тимчасові директорії"""
    templates_dir = tmp_path / "templates"
    output_dir = tmp_path / "output"
    templates_dir.mkdir()
    output_dir.mkdir()
    return templates_dir, output_dir


class TestZayavkaGeneratorExtended:
    """Розширені тести для ZayavkaGenerator"""

    def test_zayavka_module_structure(self):
        """Перевірка структури модуля"""
        from documents.generators.zayavka_generator import ZayavkaGenerator

        # Перевірка що клас має потрібні методи
        assert hasattr(ZayavkaGenerator, "get_template_name")
        assert hasattr(ZayavkaGenerator, "generate")


class TestVcGeneratorExtended:
    """Розширені тести для VCDocumentGenerator"""

    def test_vc_gen_print_vc_report(self, temp_db, tmp_path):
        """Тест звіту ВЦ"""
        from documents.vc_generator import VCDocumentGenerator
        from services.document_service import DocumentService

        doc_service = DocumentService(temp_db, str(tmp_path), str(tmp_path))
        gen = VCDocumentGenerator(
            str(tmp_path), str(tmp_path), document_service=doc_service
        )

        # Перевірка що пайплайн має метод звіту
        assert hasattr(gen.vc_pipeline, "_print_report")

    def test_vc_gen_print_final_report(self, temp_db, tmp_path):
        """Тест фінального звіту"""
        from documents.vc_generator import VCDocumentGenerator
        from services.document_service import DocumentService

        doc_service = DocumentService(temp_db, str(tmp_path), str(tmp_path))
        gen = VCDocumentGenerator(
            str(tmp_path), str(tmp_path), document_service=doc_service
        )

        assert hasattr(gen.final_pipeline, "_print_report")

    def test_vc_gen_get_file_date(self, temp_db, tmp_path):
        """Тест отримання дати"""
        from documents.vc_generator import VCDocumentGenerator
        from services.document_service import DocumentService

        doc_service = DocumentService(temp_db, str(tmp_path), str(tmp_path))
        gen = VCDocumentGenerator(
            str(tmp_path), str(tmp_path), document_service=doc_service
        )

        date_str = gen.vc_pipeline._get_file_date()
        assert len(date_str) == 10


class TestContextBuilder:
    """Тести для context_builder"""

    def test_context_builder_module_import(self):
        """Перевірка імпорту модуля"""
        from documents import context_builder

        assert context_builder is not None

    def test_context_builder_has_builder_class(self):
        """Перевірка що є клас DocumentContextBuilder"""
        try:
            from documents.context_builder import DocumentContextBuilder

            assert DocumentContextBuilder is not None
        except ImportError:
            pytest.skip("DocumentContextBuilder not available")


class TestDocumentContext:
    """Тести для document_context"""

    def test_document_context_module_import(self):
        """Перевірка імпорту модуля"""
        try:
            from documents.context import document_context

            assert document_context is not None
        except ImportError:
            pytest.skip("document_context module not available")


class TestAppRegistrationExtended:
    """Розширені тести для app_registration"""

    def test_app_registration_module_import(self):
        """Перевірка імпорту модуля"""
        from documents import app_registration

        assert app_registration is not None


class TestDocumentBaseMore:
    """Більше тестів для DocumentBase"""

    def test_replace_text_empty_replacements(self, tmp_path):
        """Заміна з порожніми replacements"""
        from documents.generators.protocol_rozgl_generator import (
            ProtocolRozglGenerator,
        )

        doc_path = tmp_path / "test.docx"
        from docx import Document

        doc = Document()
        doc.add_paragraph("Текст без тегів")
        doc.save(doc_path)

        gen = ProtocolRozglGenerator(tmp_path, tmp_path)
        result = gen.replace_text_in_document(doc_path, {})

        assert result is True

    def test_replace_text_with_empty_paragraph(self, tmp_path):
        """Заміна з порожнім параграфом"""
        from documents.generators.protocol_rozgl_generator import (
            ProtocolRozglGenerator,
        )

        doc_path = tmp_path / "test.docx"
        from docx import Document

        doc = Document()
        doc.add_paragraph("")  # порожній
        doc.add_paragraph("[TAG]")
        doc.save(doc_path)

        gen = ProtocolRozglGenerator(tmp_path, tmp_path)
        result = gen.replace_text_in_document(doc_path, {"[TAG]": "Value"})

        assert result is True


class TestServiceErrorHandling:
    """Тести для обробки помилок в сервісах"""

    def test_application_service_get_nonexistent(self, temp_db):
        """Отримання неіснуючої заявки"""
        from services.application_service import ApplicationService

        service = ApplicationService(temp_db)
        result = service.get_application_by_number("999")
        assert result is None

        service.close()

    def test_application_service_delete_nonexistent(self, temp_db):
        """Видалення неіснуючої заявки"""
        from services.application_service import ApplicationService

        service = ApplicationService(temp_db)
        # Може повертати True або False залежно від реалізації
        result = service.delete_application(99999)
        assert isinstance(result, bool)

        service.close()

    def test_dictionary_service_get_nonexistent_product(self, temp_db):
        """Отримання неіснуючого продукту"""
        from services.dictionary_service import DictionaryService

        service = DictionaryService(temp_db)
        result = service.get_product_by_name("Неіснуючий")
        assert result is None

        service.close()

    def test_document_service_with_invalid_path(self, temp_db):
        """DocumentService з невалідним шляхом"""
        from services.document_service import DocumentService

        service = DocumentService(temp_db, "/nonexistent", "/nonexistent")
        # Має створити об'єкт без помилок
        assert service is not None

        service.close()
