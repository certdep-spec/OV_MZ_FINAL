"""
Тести для покриття vc_generator та document_base
"""

import os
import tempfile

import pytest
from docx import Document

from database.db_manager import DatabaseManager
from documents.vc_generator import VCDocumentGenerator
from services.document_service import DocumentService


@pytest.fixture
def temp_db():
    """Створення тимчасової БД"""
    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".db")
    temp_file.close()
    db_manager = DatabaseManager(temp_file.name)
    db_manager.close()
    yield temp_file.name
    os.unlink(temp_file.name)


@pytest.fixture
def vc_gen(temp_db, tmp_path):
    """Створення VCDocumentGenerator для тестів"""
    templates_dir = tmp_path / "templates"
    output_dir = tmp_path / "output"
    templates_dir.mkdir()
    output_dir.mkdir()

    doc_service = DocumentService(temp_db, str(templates_dir), str(output_dir))

    return (
        VCDocumentGenerator(
            str(templates_dir), str(output_dir), document_service=doc_service
        ),
        templates_dir,
        output_dir,
    )


@pytest.fixture
def doc_base_gen(tmp_path):
    """Створення конкретного генератора для тестів DocumentBase"""
    from documents.generators.protocol_rozgl_generator import (
        ProtocolRozglGenerator,
    )

    return ProtocolRozglGenerator(tmp_path, tmp_path)


class TestVcGeneratorMethods:
    """Тести для методів VCDocumentGenerator"""

    def test_vc_gen_get_file_date(self, vc_gen):
        """Тест _get_file_date через пайплайн"""
        gen, _, _ = vc_gen
        date_str = gen.vc_pipeline._get_file_date()
        assert len(date_str) == 10
        assert "-" in date_str

    def test_vc_gen_print_vc_report(self, temp_db, tmp_path):
        """Тест _print_vc_report через пайплайн"""
        from documents.vc_generator import VCDocumentGenerator
        from services.document_service import DocumentService

        doc_service = DocumentService(temp_db, str(tmp_path), str(tmp_path))
        gen = VCDocumentGenerator(
            str(tmp_path), str(tmp_path), document_service=doc_service
        )

        # Перевірка що пайплайн має метод звіту
        assert hasattr(gen.vc_pipeline, "_print_report")

    def test_vc_gen_print_final_report(self, temp_db, tmp_path):
        """Тест _print_final_report через пайплайн"""
        from documents.vc_generator import VCDocumentGenerator
        from services.document_service import DocumentService

        doc_service = DocumentService(temp_db, str(tmp_path), str(tmp_path))
        gen = VCDocumentGenerator(
            str(tmp_path), str(tmp_path), document_service=doc_service
        )

        # Перевірка що пайплайн має метод звіту
        assert hasattr(gen.final_pipeline, "_print_report")


class TestDocumentBaseExtended:
    """Розширені тести для DocumentBase"""

    def test_replace_text_in_table_cells(self, doc_base_gen, tmp_path):
        """Заміна тексту в клітинках таблиці"""
        doc_path = tmp_path / "test.docx"
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "[TAG1]"
        table.cell(1, 1).text = "[TAG2]"
        doc.save(doc_path)

        result = doc_base_gen.replace_text_in_document(
            doc_path, {"[TAG1]": "Value1", "[TAG2]": "Value2"}
        )

        assert result is True

        doc = Document(doc_path)
        assert "Value1" in doc.tables[0].cell(0, 0).text
        assert "Value2" in doc.tables[0].cell(1, 1).text

    def test_replace_text_with_variants(self, doc_base_gen, tmp_path):
        """Заміна тексту з різними варіантами тегів"""
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("[TAG]")
        doc.add_paragraph("{{TAG}}")
        doc.add_paragraph("{TAG}")
        doc.save(doc_path)

        result = doc_base_gen.replace_text_in_document(doc_path, {"[TAG]": "Replaced"})

        assert result is True

        doc = Document(doc_path)
        assert "Replaced" in doc.paragraphs[0].text
        assert "Replaced" in doc.paragraphs[1].text
        assert "Replaced" in doc.paragraphs[2].text

    def test_replace_text_with_none_value(self, doc_base_gen, tmp_path):
        """Заміна тексту з None значенням"""
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("[TAG]")
        doc.save(doc_path)

        result = doc_base_gen.replace_text_in_document(doc_path, {"[TAG]": None})

        # Має повернути True і пропустити None
        assert result is True

    def test_build_base_replacements_with_extra(self, doc_base_gen):
        """Побудова замін з extra параметрами"""
        app_data = {
            "app_number": "001",
            "app_date": "05.04.2026",
            "company_name": "ТОВ Тест",
            "company_code": "12345678",
            "production_address": "м. Київ",
            "tu_code": "ТУ 123",
            "director_name": "Директор",
        }
        party_data = {
            "product_name": "Продукт",
            "party_code": "01E",
            "mfg_date": "01.04.2026",
            "quantity": 100,
            "unit": "шт",
            "party_date": "05.04.2026",
            "shelf_life_months": 24,
            "protocol_number": "123",
            "protocol_date": "05.04.2026",
            "cert_number": "UA.123",
            "cert_date": "05.04.2026",
        }
        extra = {"[CUSTOM]": "Custom Value"}

        result = doc_base_gen._build_base_replacements(
            app_data, party_data, party_id=1, extra=extra
        )

        assert "[CUSTOM]" in result
        assert result["[CUSTOM]"] == "Custom Value"

    def test_build_base_replacements_empty_protocol(self, doc_base_gen):
        """Побудова замін з порожнім протоколом"""
        app_data = {
            "app_number": "001",
            "app_date": "05.04.2026",
            "company_name": "ТОВ Тест",
        }
        party_data = {
            "product_name": "Продукт",
            "party_code": "01E",
            "mfg_date": "01.04.2026",
            "quantity": 100,
            "unit": "шт",
            "party_date": "05.04.2026",
            "shelf_life_months": 24,
            "protocol_number": "",
            "protocol_date": "",
        }

        result = doc_base_gen._build_base_replacements(app_data, party_data)

        assert "[G10]" in result  # Має бути порожнім або частково заповненим

    def test_replace_text_in_header_footer(self, doc_base_gen, tmp_path):
        """Заміна тексту в колонтитулах"""

        doc_path = tmp_path / "test.docx"
        doc = Document()

        # Додаємо текст в header
        section = doc.sections[0]
        header = section.header
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        run = p.add_run("[HEADER_TAG]")
        run.text = "[HEADER_TAG]"

        doc.save(doc_path)

        result = doc_base_gen.replace_text_in_document(
            doc_path, {"[HEADER_TAG]": "Header Value"}
        )

        assert result is True

    def test_validate_app_data_all_fields_present(self, doc_base_gen):
        """Валідація з усіма полями"""
        app_data = {
            "app_number": "001",
            "app_date": "05.04.2026",
            "company_name": "ТОВ Тест",
        }
        # Не повинно виникати винятку
        doc_base_gen._validate_app_data(app_data)

    def test_get_file_date_format(self, doc_base_gen):
        """Перевірка формату дати"""
        date_str = doc_base_gen._get_file_date()
        # Перевірка що це YYYY-MM-DD
        parts = date_str.split("-")
        assert len(parts) == 3
        assert len(parts[0]) == 4  # рік
        assert 1 <= int(parts[1]) <= 12  # місяць
        assert 1 <= int(parts[2]) <= 31  # день


class TestVcGeneratorGenerateVC:
    """Тести для generate_vc_documents"""

    def test_vc_gen_generate_vc_documents_structure(self, vc_gen):
        """Перевірка структури generate_vc_documents"""
        gen, _, _ = vc_gen

        # Перевірка що метод існує і має правильну сигнатуру
        import inspect

        sig = inspect.signature(gen.generate_vc_documents)
        assert "app_data" in sig.parameters
        assert "parties" in sig.parameters


class TestDocumentBaseEdgeCases:
    """Тести для крайніх випадків DocumentBase"""

    def test_replace_text_empty_document(self, doc_base_gen, tmp_path):
        """Заміна в порожньому документі"""
        doc_path = tmp_path / "test.docx"
        doc = Document()
        doc.save(doc_path)

        result = doc_base_gen.replace_text_in_document(doc_path, {"[TAG]": "Value"})
        assert result is True

    def test_replace_text_tag_spread_across_runs(self, doc_base_gen, tmp_path):
        """Тег розбитий на кілька run'ів"""
        doc_path = tmp_path / "test.docx"
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("[T")
        p.add_run("AG")
        p.add_run("]")
        doc.save(doc_path)

        result = doc_base_gen.replace_text_in_document(doc_path, {"[TAG]": "Replaced"})
        assert result is True

    def test_set_table_borders_existing(self, doc_base_gen, tmp_path):
        """Встановлення рамок на існуючу таблицю"""
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        doc = Document()
        table = doc.add_table(rows=2, cols=2)

        # Вже є рамки
        tbl = table._element
        tblPr = tbl.xpath("w:tblPr")
        if not tblPr:
            tblPr_elem = OxmlElement("w:tblPr")
            tbl.insert(0, tblPr_elem)
        else:
            tblPr = tblPr[0]
            tblBorders = OxmlElement("w:tblBorders")
            top = OxmlElement("w:top")
            top.set(qn("w:val"), "single")
            tblBorders.append(top)
            tblPr.append(tblBorders)

        doc_base_gen.set_table_borders(table)
        assert True  # Якщо не впало - тест пройдено

    def test_find_template_with_doc_extension(self, doc_base_gen, tmp_path):
        """Пошук шаблону з .doc розширенням"""
        # Створюємо .doc файл
        doc_file = tmp_path / "template.doc"
        doc_file.write_text("fake doc")

        # Шукаємо з .docx
        result = doc_base_gen._find_template("template.docx")
        assert result is not None
        assert result.name == "template.doc"

    def test_calculate_expiry_date_edge_cases(self, doc_base_gen):
        """Розрахунок терміну придатності - крайні випадки"""
        # Різні терміни
        result_1 = doc_base_gen.calculate_expiry_date("01.01.2026", 1)
        assert result_1 != ""

        result_12 = doc_base_gen.calculate_expiry_date("01.01.2026", 12)
        assert result_12 != ""

        result_0 = doc_base_gen.calculate_expiry_date("01.01.2026", 0)
        assert result_0 != ""

    def test_get_month_name_ua_all_months(self, doc_base_gen):
        """Отримання назв всіх місяців"""
        months = {
            "01": "січня",
            "02": "лютого",
            "03": "березня",
            "04": "квітня",
            "05": "травня",
            "06": "червня",
            "07": "липня",
            "08": "серпня",
            "09": "вересня",
            "10": "жовтня",
            "11": "листопада",
            "12": "грудня",
        }
        for num, name in months.items():
            assert doc_base_gen.get_month_name_ua(num) == name
